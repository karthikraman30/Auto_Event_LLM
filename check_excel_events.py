#!/usr/bin/env python3
"""Check events from Excel file"""
import openpyxl
from datetime import datetime, timedelta

# Load the Excel file
wb = openpyxl.load_workbook('event_category/events.xlsx')
ws = wb.active

# Calculate date range
today = datetime.now().date()
one_month_later = today + timedelta(days=30)

total_events = ws.max_row - 1  # Subtract header row

print(f"{'='*80}")
print(f"EVENTS SCRAPED FROM: https://biblioteket.stockholm.se/evenemang")
print(f"{'='*80}\n")
print(f"Total events scraped: {total_events}")
print(f"Date range filter: {today} to {one_month_later}\n")

# Count events in next 1 month (column 3 is date_iso)
events_in_range = 0
for row in range(2, ws.max_row + 1):
    date_iso = ws.cell(row=row, column=3).value
    if date_iso:
        try:
            event_date = datetime.strptime(str(date_iso), "%Y-%m-%d").date()
            if today <= event_date <= one_month_later:
                events_in_range += 1
        except:
            pass

print(f"Events in next 1 month: {events_in_range}")
print(f"\n{'='*80}")
print("Sample events (first 10):")
print(f"{'='*80}\n")

# Show first 10 events
for i in range(min(10, ws.max_row - 1)):
    row = i + 2
    name = ws.cell(row=row, column=1).value
    date = ws.cell(row=row, column=3).value
    time = ws.cell(row=row, column=5).value
    location = ws.cell(row=row, column=6).value
    target = ws.cell(row=row, column=8).value
    
    print(f"{i+1}. {name}")
    print(f"   Date: {date} | Time: {time}")
    print(f"   Location: {location} | Target: {target}")
    print()

# Show date distribution
print(f"{'='*80}")
print("Date distribution (all dates):")
print(f"{'='*80}\n")

date_counts = {}
for row in range(2, ws.max_row + 1):
    date_iso = ws.cell(row=row, column=3).value
    if date_iso:
        date_counts[str(date_iso)] = date_counts.get(str(date_iso), 0) + 1

for date in sorted(date_counts.keys())[:20]:  # Show first 20 dates
    print(f"{date}: {date_counts[date]} events")
